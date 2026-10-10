"""Monthly dashboard: route, metrics (computed in SQL), charts data, accessibility and empty state."""

import calendar
import json
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

import dashboard as dashboard_module
import utils
from extensions import db
from models import Item, User

CHART_URL = "https://cdn.jsdelivr.net/npm/chart.js@4.5.1/dist/chart.umd.min.js"
EMOJI_PATTERN = re.compile("[\u2600-\u27bf\u2b00-\u2bff\U0001f000-\U0001faff\ufe0f]")
ENGLISH_WORDS = re.compile(
    r"\b(Dashboard|Lost|Found|Returned|Month|Reports?|Export|Previous|Next|Category|Location|"
    r"Status|Date|Items?|Actions?|Price|Quantity|Stock|Loading|Overview|Latest|Recent|Show|Filter)\b"
)
EMPTY_TEXT = "No hay reportes con fecha del suceso en este mes."


# --- helpers ------------------------------------------------------------------


def _text(response):
    return response.get_data(as_text=True)


def _main(html):
    return html.split("<main", 1)[1].split("</main>", 1)[0]


def _visible(html):
    """Text a person reads: no scripts, styles or tags, plus accessible names and tooltips."""
    attributes = re.findall(r'(?:aria-label|title|alt|placeholder)="([^"]*)"', html)
    stripped = re.sub(r"<(script|style)\b.*?</\1>", " ", html, flags=re.DOTALL)
    return re.sub(r"<[^>]+>", " ", stripped) + " " + " ".join(attributes)


def _kpi(html, name):
    match = re.search(rf'data-kpi="{name}"[^>]*>\s*(.*?)\s*</', html, re.DOTALL)
    assert match is not None, f"no KPI named {name}"
    return match.group(1)


def _chart_data(html):
    match = re.search(
        r'<script type="application/json" id="dashboard-chart-data">(.*?)</script>',
        html,
        re.DOTALL,
    )
    assert match is not None, "chart data block not found"
    return json.loads(match.group(1))


def _nav_href(html, rel):
    for tag in re.findall(r"<a\b[^>]*>", html):
        if re.search(rf'rel="{rel}"', tag):
            return re.search(r'href="([^"]+)"', tag).group(1)
    return None


def _recent_ids(html):
    section = re.search(r'<table[^>]*id="dashboard-recent"[^>]*>(.*?)</table>', html, re.DOTALL)
    assert section is not None, "recent reports table not found"
    return [int(found) for found in re.findall(r'href="/item/(\d+)"', section.group(1))]


def _table_rows(html, table_id):
    section = re.search(rf'<table[^>]*id="{table_id}"[^>]*>(.*?)</table>', html, re.DOTALL)
    assert section is not None, f"table {table_id} not found"
    body = section.group(1).split("<tbody", 1)[1]
    rows = []
    for row in re.findall(r"<tr\b.*?</tr>", body, re.DOTALL):
        cells = re.findall(r"<t[dh]\b[^>]*>(.*?)</t[dh]>", row, re.DOTALL)
        rows.append([re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", cell)).strip() for cell in cells])
    return rows


def _add_reports(owner, *specs):
    """Persist reports in bulk. Each spec: dict with day (ISO text, or None) plus Item fields."""
    items = []
    for spec in specs:
        spec = dict(spec)
        day = spec.pop("day", None)
        values = {
            "name": "Objeto",
            "location": "Biblioteca central",
            "category": None,
            "status": "lost",
            "event_date": date.fromisoformat(day) if day else None,
            "approval": "approved",
        }
        values.update(spec)
        items.append(Item(user_id=owner.id, **values))
    db.session.add_all(items)
    db.session.commit()
    return items


def _month_of(count, day_start=1, **fields):
    return [dict(day=f"2026-10-{day_start + (i % 28):02d}", **fields) for i in range(count)]


@pytest.fixture
def maria(browser_for):
    """The administrator: the panel, import and export belong to that role."""
    browser = browser_for("maria")
    user = User.query.filter_by(username="maria").one()
    user.is_admin = True
    db.session.commit()
    return browser


@pytest.fixture
def maria_user(maria):
    return User.query.filter_by(username="maria").first()


@pytest.fixture
def carlos_user(maria, make_user):
    return make_user("carlos", "password123")


@pytest.fixture
def freeze_now(monkeypatch):
    """Freeze the clock seen by utils.today_local (optionally in a fixed-offset display zone)."""

    def _freeze(year, month, day, hour=12, offset_hours=None):
        base = datetime(year, month, day, hour, 0, tzinfo=timezone.utc)

        class FrozenDatetime(datetime):
            @classmethod
            def now(cls, tz=None):
                return base.astimezone(tz) if tz else base.replace(tzinfo=None)

        monkeypatch.setattr(utils, "datetime", FrozenDatetime)
        if offset_hours is not None:
            zone = timezone(timedelta(hours=offset_hours))
            monkeypatch.setattr(utils, "_display_zone", lambda: zone)

    return _freeze


def _get(browser, month="2026-10"):
    response = browser.get(f"/dashboard/?mes={month}")
    assert response.status_code == 200
    return _text(response)


# --- routes -------------------------------------------------------------------


def test_dashboard_requires_login(client):
    response = client.get("/dashboard/")

    assert response.status_code == 302
    location = urlsplit(response.headers["Location"])
    assert location.path == "/auth/login"
    assert parse_qs(location.query)["next"] == ["/dashboard/"]


def test_anonymous_legacy_url_ends_at_the_login_page(client):
    response = client.get("/dashboard?mes=2026-02", follow_redirects=True)

    assert response.request.path == "/auth/login"
    assert "Iniciar sesión" in _text(response)


def test_legacy_url_redirects_to_the_single_dashboard_keeping_the_query(maria):
    response = maria.get("/dashboard?mes=2026-02")

    assert response.status_code == 302
    location = urlsplit(response.headers["Location"])
    assert location.path == "/dashboard/"
    assert location.query == "mes=2026-02"


def test_legacy_url_without_query_has_no_trailing_question_mark(maria):
    response = maria.get("/dashboard")

    assert response.status_code == 302
    assert urlsplit(response.headers["Location"]).path == "/dashboard/"
    assert "?" not in response.headers["Location"]


def test_legacy_url_keeps_every_query_parameter_untouched(maria):
    response = maria.get("/dashboard?mes=2026-02&x=a%20b")

    assert urlsplit(response.headers["Location"]).query == "mes=2026-02&x=a%20b"


def test_legacy_url_redirect_can_be_followed(maria):
    response = maria.get("/dashboard?mes=2026-02", follow_redirects=True)

    assert response.request.path == "/dashboard/"
    assert "febrero de 2026" in _text(response)


def test_dashboard_route_table(app):
    rules = {
        rule.rule: rule.endpoint
        for rule in app.url_map.iter_rules()
        if rule.endpoint.startswith("dashboard.") or rule.rule.startswith("/dashboard")
    }

    assert rules["/dashboard/"] == "dashboard.home"
    assert rules["/dashboard"] == "dashboard.legacy_redirect"
    assert "main.dashboard" not in app.view_functions
    assert len(rules) == 2


def test_dashboard_only_accepts_get(maria):
    assert maria.post("/dashboard/").status_code == 405
    assert maria.post("/dashboard").status_code == 405


def test_navbar_marks_the_panel_link_as_current(maria):
    html = _text(maria.get("/dashboard/"))

    active = re.findall(r'<a[^>]*aria-current="page"[^>]*>', html)
    assert len(active) == 1
    assert 'href="/dashboard/"' in active[0]


# --- month selection ----------------------------------------------------------


def test_default_month_is_the_current_month(maria, freeze_now):
    freeze_now(2026, 10, 15)

    html = _text(maria.get("/dashboard/"))

    assert "octubre de 2026" in html
    assert re.search(r'<input[^>]*type="month"[^>]*value="2026-10"', html)


def test_default_month_follows_the_display_timezone(maria, freeze_now):
    # 2026-11-01 02:00 UTC is still 2026-10-31 21:00 at UTC-5.
    freeze_now(2026, 11, 1, hour=2, offset_hours=-5)

    assert "octubre de 2026" in _text(maria.get("/dashboard/"))


def test_default_month_uses_utc_when_no_other_zone_is_configured(maria, freeze_now):
    freeze_now(2026, 11, 1, hour=2)

    assert "noviembre de 2026" in _text(maria.get("/dashboard/"))


@pytest.mark.parametrize(
    "value",
    ["2026-13", "2026-00", "abc", "2026-1", "202610", "", "0000-01", "9999-01", "2026-10-01", "2026-10; DROP"],
)
def test_invalid_month_falls_back_to_the_current_month(maria, freeze_now, value):
    freeze_now(2026, 7, 4)

    html = _text(maria.get("/dashboard/", query_string={"mes": value}))

    assert "julio de 2026" in html
    assert re.search(r'<input[^>]*type="month"[^>]*value="2026-07"', html)


def test_explicit_month_is_used(maria, freeze_now):
    freeze_now(2026, 10, 15)

    html = _get(maria, "2025-03")

    assert "marzo de 2025" in html
    assert re.search(r'<input[^>]*type="month"[^>]*value="2025-03"', html)


def test_header_is_the_spanish_monthly_panel(maria):
    html = _get(maria, "2026-10")

    assert "Panel mensual" in html
    assert re.search(r"<title>\s*Panel mensual · Objetos perdidos\s*</title>", html)
    assert '<p class="page-subtitle">octubre de 2026</p>' in html


@pytest.mark.parametrize(
    "month, previous, following",
    [
        ("2026-10", "2026-09", "2026-11"),
        ("2026-12", "2026-11", "2027-01"),
        ("2026-01", "2025-12", "2026-02"),
        ("2024-02", "2024-01", "2024-03"),
    ],
)
def test_month_navigation_links(maria, month, previous, following):
    html = _get(maria, month)

    assert _nav_href(html, "prev") == f"/dashboard/?mes={previous}"
    assert _nav_href(html, "next") == f"/dashboard/?mes={following}"
    assert "Mes anterior" in html and "Mes siguiente" in html


def test_navigation_stops_at_the_supported_years(maria):
    first = _get(maria, "0001-01")
    last = _get(maria, "9998-12")

    assert _nav_href(first, "prev") is None
    assert _nav_href(first, "next") == "/dashboard/?mes=0001-02"
    assert _nav_href(last, "next") is None
    assert _nav_href(last, "prev") == "/dashboard/?mes=9998-11"
    assert 'aria-disabled="true"' in first and 'aria-disabled="true"' in last


def test_month_form_submits_mes_with_get_and_a_visible_button(maria):
    html = _get(maria, "2026-10")

    form = re.search(r'<form[^>]*id="month-form"[^>]*>(.*?)</form>', html, re.DOTALL)
    assert form is not None
    assert re.search(r'method="get"', form.group(0), re.IGNORECASE)
    assert 'action="/dashboard/"' in form.group(0)
    assert re.search(r'<input[^>]*type="month"[^>]*name="mes"[^>]*value="2026-10"', form.group(1))
    assert re.search(r'<label[^>]*for="dashboard-month"', form.group(1))
    assert "Ver mes" in form.group(1)
    assert "csrf_token" not in form.group(1)


def test_export_button_carries_the_selected_month(maria):
    html = _get(maria, "2026-02")

    link = re.search(r'<a[^>]*href="/export/excel\?mes=2026-02"[^>]*>(.*?)</a>', html, re.DOTALL)
    assert link is not None
    assert "Exportar mes (Excel)" in link.group(1)
    assert "btn-primary" in link.group(0)


def test_export_link_downloads_the_monthly_file(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05", name="Paraguas"), dict(day="2026-09-05", name="Otro mes"))
    html = _get(maria, "2026-10")
    href = re.search(r'href="(/export/excel\?mes=2026-10)"', html).group(1)

    response = maria.get(href)

    assert response.headers["Content-Disposition"] == "attachment; filename=objetos_2026-10.xlsx"
    sheet = load_workbook(BytesIO(response.get_data()))["Reportes"]
    names = [row[2] for row in sheet.iter_rows(min_row=5, values_only=True)]
    assert "Paraguas" in names and "Otro mes" not in names


# --- KPIs ---------------------------------------------------------------------


def test_kpi_counts_per_status(maria, maria_user):
    _add_reports(
        maria_user,
        *[dict(day="2026-10-02", status="lost")] * 3,
        *[dict(day="2026-10-03", status="found")] * 2,
        dict(day="2026-10-04", status="returned"),
    )

    html = _get(maria)

    assert _kpi(html, "total") == "6"
    assert _kpi(html, "lost") == "3"
    assert _kpi(html, "found") == "2"
    assert _kpi(html, "returned") == "1"
    assert _kpi(html, "return_rate") == "16,7 %"


def test_kpi_labels_and_notes_are_spanish(maria):
    html = _main(_get(maria))

    for text in (
        "Objetos reportados",
        "Perdidos o encontrados con fecha del suceso en el mes",
        "Perdidos",
        "Encontrados",
        "Entregados",
    ):
        assert text in html, text
    assert html.count("kpi-card--primary") == 1
    assert re.search(r'kpi-card--primary[^"]*"[^>]*>\s*<p class="kpi-label">Objetos reportados', html)


def test_five_kpi_cards_are_rendered_even_with_zero_reports(maria):
    html = _get(maria, "2026-10")

    assert len(re.findall(r'<article class="card kpi-card', html)) == 5
    for name in ("total", "lost", "found", "returned"):
        assert _kpi(html, name) == "0"
    assert _kpi(html, "return_rate") == "—"


@pytest.mark.parametrize(
    "returned, others, expected",
    [
        (0, 0, "—"),
        (3, 4, "42,9 %"),
        (0, 4, "0,0 %"),
        (5, 0, "100,0 %"),
        (1, 2, "33,3 %"),
        (2, 1, "66,7 %"),
    ],
)
def test_return_rate_formatting(maria, maria_user, returned, others, expected):
    _add_reports(
        maria_user,
        *[dict(day="2026-10-10", status="returned")] * returned,
        *[dict(day="2026-10-11", status="lost")] * others,
    )

    assert _kpi(_get(maria), "return_rate") == expected


@pytest.mark.parametrize(
    "current, previous, expected",
    [
        (5, 2, "+3 respecto al mes anterior"),
        (2, 4, "-2 respecto al mes anterior"),
        (3, 3, "Sin cambios respecto al mes anterior"),
        (0, 0, "Sin datos del mes anterior"),
        (4, 0, "+4 respecto al mes anterior"),
        (0, 3, "-3 respecto al mes anterior"),
        (1, 0, "+1 respecto al mes anterior"),
    ],
)
def test_delta_versus_previous_month(maria, maria_user, current, previous, expected):
    _add_reports(
        maria_user,
        *[dict(day="2026-10-10")] * current,
        *[dict(day="2026-09-30")] * previous,
    )

    assert _kpi(_get(maria, "2026-10"), "delta") == expected


def test_delta_of_january_compares_with_the_previous_december(maria, maria_user):
    _add_reports(
        maria_user,
        *[dict(day="2026-01-15")] * 4,
        *[dict(day="2025-12-31")] * 1,
        *[dict(day="2026-02-01")] * 9,
    )

    assert _kpi(_get(maria, "2026-01"), "delta") == "+3 respecto al mes anterior"


def test_delta_uses_thousands_separator(maria, maria_user):
    _add_reports(maria_user, *[dict(day="2026-10-10")] * 1203)

    assert _kpi(_get(maria, "2026-10"), "total") == "1.203"
    assert _kpi(_get(maria, "2026-10"), "delta") == "+1.203 respecto al mes anterior"


def test_reports_of_every_user_are_counted(maria, maria_user, carlos_user):
    _add_reports(maria_user, dict(day="2026-10-02"), dict(day="2026-10-03", status="found"))
    _add_reports(carlos_user, dict(day="2026-10-04"), dict(day="2026-10-05", status="returned"))

    html = _get(maria)

    assert _kpi(html, "total") == "4"
    assert _kpi(html, "lost") == "2"
    assert _kpi(html, "found") == "1"
    assert _kpi(html, "returned") == "1"
    assert sorted(_recent_ids(html)) == sorted(item.id for item in Item.query.all())


def test_regular_users_cannot_open_the_panel(browser_for, maria_user, carlos_user):
    """Only the administrator has the monthly panel; others land on the report center."""
    response = browser_for("carlos").get("/dashboard/?mes=2026-10")

    assert response.status_code == 302
    assert response.headers["Location"].split("?", 1)[0] == "/"


def test_only_reports_inside_the_month_are_counted(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-09-30"),
        dict(day="2026-10-01"),
        dict(day="2026-10-31"),
        dict(day="2026-11-01"),
        dict(day="2025-10-15"),
    )

    html = _get(maria)

    assert _kpi(html, "total") == "2"
    assert _kpi(html, "delta") == "+1 respecto al mes anterior"


def test_reports_without_event_date_are_ignored(maria, maria_user):
    _add_reports(maria_user, dict(day=None), dict(day=None, status="returned"), dict(day="2026-10-10"))

    html = _get(maria)

    assert _kpi(html, "total") == "1"
    assert _kpi(html, "returned") == "0"
    assert sum(_chart_data(html)["daily"]["counts"]) == 1
    assert len(_recent_ids(html)) == 1


def test_unknown_status_counts_in_the_total_only(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-10", status="archived"), dict(day="2026-10-10", status="lost"))

    html = _get(maria)

    assert _kpi(html, "total") == "2"
    assert _kpi(html, "lost") == "1"
    assert _kpi(html, "found") == "0"
    assert _kpi(html, "returned") == "0"


# --- per-day series -----------------------------------------------------------


@pytest.mark.parametrize(
    "month, days",
    [
        ("2026-02", 28),
        ("2028-02", 29),
        ("2026-04", 30),
        ("2026-10", 31),
        ("2026-12", 31),
        ("2026-01", 31),
        ("1900-02", 28),
        ("2000-02", 29),
    ],
)
def test_daily_series_has_one_zero_filled_entry_per_day(app, month, days):
    year, number = (int(value) for value in month.split("-"))
    data = dashboard_module.month_summary(year, number)["daily"]

    assert data["labels"] == [str(day) for day in range(1, days + 1)]
    assert data["counts"] == [0] * days


def test_daily_series_places_counts_on_their_day(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2028-02-01"),
        dict(day="2028-02-29"),
        dict(day="2028-02-29", status="found"),
        dict(day="2028-03-01"),
        dict(day="2028-01-31"),
    )

    counts = _chart_data(_get(maria, "2028-02"))["daily"]["counts"]

    assert len(counts) == 29
    assert counts[0] == 1
    assert counts[28] == 2
    assert sum(counts) == 3


def test_daily_series_matches_the_calendar_for_every_month_of_a_year(app, maria_user):
    with app.test_request_context("/"):
        for month in range(1, 13):
            summary = dashboard_module.month_summary(2026, month)
            assert len(summary["daily"]["counts"]) == calendar.monthrange(2026, month)[1]


def test_daily_total_equals_the_total_kpi(maria, maria_user):
    _add_reports(maria_user, *_month_of(37, status="found"))

    html = _get(maria)

    assert sum(_chart_data(html)["daily"]["counts"]) == int(_kpi(html, "total")) == 37


# --- categories ---------------------------------------------------------------


def test_categories_use_spanish_labels_ordered_by_count(maria, maria_user):
    _add_reports(
        maria_user,
        *[dict(day="2026-10-02", category="keys")] * 3,
        *[dict(day="2026-10-02", category="documents")] * 5,
        dict(day="2026-10-02", category="bags"),
        *[dict(day="2026-10-02", category="electronics")] * 2,
    )

    data = _chart_data(_get(maria))["categories"]

    assert data["labels"] == ["Documentos", "Llaves", "Electrónica", "Mochilas y útiles"]
    assert data["counts"] == [5, 3, 2, 1]


def test_categories_handle_missing_legacy_and_free_text_values(maria, maria_user):
    _add_reports(
        maria_user,
        *[dict(day="2026-10-02", category=None)] * 2,
        dict(day="2026-10-02", category=""),
        dict(day="2026-10-02", category="   "),
        *[dict(day="2026-10-02", category="general")] * 2,
        dict(day="2026-10-02", category="Herramientas"),
    )

    data = _chart_data(_get(maria))["categories"]

    assert dict(zip(data["labels"], data["counts"])) == {
        "Sin categoría": 4,
        "general": 2,
        "Herramientas": 1,
    }
    assert data["labels"][0] == "Sin categoría"
    assert len(set(data["labels"])) == len(data["labels"])


def test_categories_with_the_same_count_are_ordered_by_label(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-02", category="other"),
        dict(day="2026-10-02", category="keys"),
        dict(day="2026-10-02", category="documents"),
    )

    assert _chart_data(_get(maria))["categories"]["labels"] == ["Documentos", "Llaves", "Otros"]


def test_categories_ignore_reports_of_other_months(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-02", category="keys"), dict(day="2026-09-02", category="bags"))

    assert _chart_data(_get(maria))["categories"]["labels"] == ["Llaves"]


# --- locations ----------------------------------------------------------------


def test_locations_are_capped_at_five_and_ordered(maria, maria_user):
    specs = []
    for place, amount in (("Cafetería", 2), ("Biblioteca", 6), ("Aula 4", 3), ("Gimnasio", 1), ("Laboratorio", 4), ("Patio", 5), ("Auditorio", 1)):
        specs += [dict(day="2026-10-02", location=place)] * amount
    _add_reports(maria_user, *specs)

    data = _chart_data(_get(maria))["locations"]

    assert data["labels"] == ["Biblioteca", "Patio", "Laboratorio", "Aula 4", "Cafetería"]
    assert data["counts"] == [6, 5, 4, 3, 2]


def test_locations_with_the_same_count_are_ordered_by_name(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-02", location="Patio"),
        dict(day="2026-10-02", location="Biblioteca"),
        dict(day="2026-10-02", location="Auditorio"),
    )

    assert _chart_data(_get(maria))["locations"]["labels"] == ["Auditorio", "Biblioteca", "Patio"]


def test_missing_and_blank_locations_are_grouped_as_unspecified(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-02", location=None),
        dict(day="2026-10-02", location=""),
        dict(day="2026-10-02", location="   "),
        dict(day="2026-10-02", location="Patio"),
    )

    data = _chart_data(_get(maria))["locations"]

    assert data["labels"] == ["Sin especificar", "Patio"]
    assert data["counts"] == [3, 1]


def test_locations_are_trimmed_before_grouping(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-02", location="Patio"), dict(day="2026-10-02", location="  Patio "))

    data = _chart_data(_get(maria))["locations"]

    assert data["labels"] == ["Patio"] and data["counts"] == [2]


# --- recent reports -----------------------------------------------------------


def test_recent_table_is_limited_to_ten_and_ordered(maria, maria_user):
    items = _add_reports(
        maria_user,
        *[dict(day=f"2026-10-{day:02d}", name=f"Objeto {day:02d}") for day in range(1, 13)],
    )
    twin = _add_reports(maria_user, dict(day="2026-10-12", name="Gemelo"))[0]

    ids = _recent_ids(_get(maria))

    by_day = {item.event_date.day: item.id for item in items}
    assert len(ids) == 10
    # Newest event date first; on the same date the highest id comes first.
    assert ids == [twin.id, by_day[12], by_day[11], by_day[10], by_day[9], by_day[8], by_day[7], by_day[6], by_day[5], by_day[4]]


def test_recent_table_columns_and_content(maria, maria_user):
    item = _add_reports(
        maria_user,
        dict(day="2026-10-05", name="Mochila azul", location="Aula 4", category="bags", status="found"),
    )[0]

    html = _get(maria)

    rows = _table_rows(html, "dashboard-recent")
    assert rows == [["Mochila azul", "Aula 4", "Mochilas y útiles", "5 de octubre de 2026", "Encontrado"]]
    header = re.search(r'id="dashboard-recent".*?<thead>(.*?)</thead>', html, re.DOTALL).group(1)
    for column in ("Objeto", "Lugar", "Categoría", "Fecha", "Estado"):
        assert column in header
    assert f'href="/item/{item.id}"' in html
    assert "badge-info" in html


def test_recent_table_handles_missing_values_and_escapes_html(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05", name="<b>negrita</b>", location=None, category=None, status="returned"))

    html = _get(maria)

    assert "<b>negrita</b>" not in html
    assert "&lt;b&gt;negrita&lt;/b&gt;" in html
    assert _table_rows(html, "dashboard-recent")[0][1:3] == ["Sin especificar", "Sin categoría"]
    assert "badge-success" in html


def test_link_to_all_reports_of_the_month(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    html = _get(maria, "2026-10")

    link = re.search(r'<a[^>]*href="(/\?mes=2026-10)"[^>]*>(.*?)</a>', html, re.DOTALL)
    assert link is not None
    assert "Ver todos los reportes del mes" in link.group(2)


# --- chart data, scripts and accessibility ------------------------------------


def test_chart_data_is_valid_json_with_aligned_series(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-05", category="keys", location="Patio"),
        dict(day="2026-10-06", category="bags", location="Aula 4"),
    )

    data = _chart_data(_get(maria))

    assert data["month"] == "octubre de 2026"
    for key in ("daily", "categories", "locations"):
        assert set(data[key]) == {"labels", "counts"}
        assert len(data[key]["labels"]) == len(data[key]["counts"])
        assert all(isinstance(count, int) for count in data[key]["counts"])
        assert all(isinstance(label, str) for label in data[key]["labels"])


def test_chart_data_cannot_break_out_of_the_script_block(maria, maria_user):
    hostile = "</script><script>alert(1)</script>"
    _add_reports(maria_user, dict(day="2026-10-05", location=hostile, category="<img src=x onerror=alert(1)>"))

    html = _get(maria)

    assert "<script>alert(1)" not in html
    assert "<img src=x" not in html
    data = _chart_data(html)
    assert data["locations"]["labels"] == [hostile]
    assert data["categories"]["labels"] == ["<img src=x onerror=alert(1)>"]


def test_chart_library_is_pinned_and_loaded_only_when_charts_are_shown(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    html = _get(maria)

    assert html.count(CHART_URL) == 1
    assert len(re.findall(r"<script[^>]*chart\.js", html)) == 1
    assert "chart.js@4.4.7" not in html and "chart.js\"" not in html
    # The library must load before the script that draws the charts.
    assert html.index(CHART_URL) < html.index('getElementById("dashboard-chart-data")')


def test_chart_library_is_not_part_of_the_base_layout(maria, browser_for):
    base = (Path(dashboard_module.__file__).parent / "templates" / "base.html").read_text(encoding="utf-8")

    assert "chart.js" not in base.lower()
    assert "chart.js" not in _text(browser_for().get("/auth/login")).lower()
    assert "chart.js" not in _text(maria.get("/")).lower()
    assert "chart.js" not in _text(maria.get("/import/csv")).lower()


def test_drawing_script_survives_a_blocked_cdn(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    html = _get(maria)

    assert 'typeof Chart === "undefined"' in html
    assert "ticks" in html and "precision: 0" in html
    assert "beginAtZero: true" in html
    assert "responsive: true" in html
    assert "display: false" in html  # no legend on single-dataset charts
    assert 'indexAxis: "y"' in html
    assert "1 reporte" in html and "reportes" in html


def test_chart_cards_with_accessible_canvases(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    html = _get(maria)

    for title in ("Reportes por día", "Reportes por categoría", "Lugares con más reportes", "Estado de los reportes"):
        assert title in html, title
    canvases = re.findall(r"<canvas\b[^>]*>", html)
    assert len(canvases) == 4
    for canvas in canvases:
        assert re.search(r'role="img"', canvas)
        label = re.search(r'aria-label="([^"]+)"', canvas)
        assert label is not None and len(label.group(1)) > 10
    assert html.count("chart-card") >= 3


def test_every_chart_has_a_data_table_fallback(maria, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-05", category="keys", location="Patio"),
        dict(day="2026-10-05", category="keys", location="Patio"),
        dict(day="2026-10-06", category="bags", location="Aula 4"),
    )

    html = _get(maria)
    data = _chart_data(html)

    assert html.count("<details") >= 3
    assert "Ver datos en tabla" in html
    daily = _table_rows(html, "chart-table-daily")
    assert [row[0] for row in daily] == data["daily"]["labels"]
    assert [int(row[1]) for row in daily] == data["daily"]["counts"]
    assert _table_rows(html, "chart-table-categories") == [["Llaves", "2"], ["Mochilas y útiles", "1"]]
    assert _table_rows(html, "chart-table-locations") == [["Patio", "2"], ["Aula 4", "1"]]


# --- empty state --------------------------------------------------------------


def test_empty_month_shows_the_empty_state_and_hides_charts(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-09-05"))

    html = _get(maria, "2026-10")
    body = _main(html)

    assert EMPTY_TEXT in body
    assert "empty-state" in body
    # The sidebar already offers "Nuevo reporte"; the empty state adds no duplicate buttons.
    assert "Reportar objeto" not in body and "Ver otro mes" not in body
    assert "<canvas" not in body
    assert "dashboard-chart-data" not in html
    assert "chart-card" not in body
    assert "dashboard-recent" not in body
    assert CHART_URL not in html
    # Navigation, export and KPI cards stay available.
    assert _nav_href(html, "prev") == "/dashboard/?mes=2026-09"
    assert "Exportar mes (Excel)" in body
    assert 'type="month"' in body
    assert _kpi(html, "total") == "0"
    assert _kpi(html, "delta") == "-1 respecto al mes anterior"


def test_month_with_reports_has_no_empty_state(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    assert EMPTY_TEXT not in _get(maria)


def test_empty_database_renders(maria):
    response = maria.get("/dashboard/")

    assert response.status_code == 200
    assert EMPTY_TEXT in _text(response)
    assert "<canvas" not in _text(response)
    assert CHART_URL not in _text(response)
    assert "dashboard-chart-data" not in _text(response)


# --- formal Spanish, no emojis ------------------------------------------------


@pytest.mark.parametrize("with_data", [True, False])
def test_page_is_spanish_and_has_no_emojis(maria, maria_user, with_data):
    if with_data:
        _add_reports(
            maria_user,
            dict(day="2026-10-05", category="keys", location="Patio"),
            dict(day="2026-10-06", category=None, location=None, status="returned"),
        )

    html = _get(maria)
    visible = _visible(_main(html))

    assert '<html lang="es"' in html
    assert not EMOJI_PATTERN.search(html)
    assert not ENGLISH_WORDS.search(visible), ENGLISH_WORDS.search(visible)
    for leftover in ("Inventario", "Precio", "Cantidad", "$", "stock", "Stock"):
        assert leftover not in visible, leftover


def test_sentence_case_headings(maria, maria_user):
    _add_reports(maria_user, dict(day="2026-10-05"))

    html = _main(_get(maria))

    for heading in ("Reportes por día", "Reportes por categoría", "Lugares con más reportes", "Últimos reportes"):
        assert heading in html
    for title_case in ("Reportes Por", "Lugares Con", "Panel Mensual", "Objetos Reportados"):
        assert title_case not in html


# --- performance sanity -------------------------------------------------------


class _QueryCounter:
    def __init__(self):
        self.statements = []

    def __call__(self, conn, cursor, statement, parameters, context, executemany):
        self.statements.append(statement)


def _count_queries(browser, url):
    counter = _QueryCounter()
    event.listen(db.engine, "before_cursor_execute", counter)
    try:
        response = browser.get(url)
    finally:
        event.remove(db.engine, "before_cursor_execute", counter)
    assert response.status_code == 200
    return counter.statements


def test_dashboard_uses_a_small_bounded_number_of_queries(maria, maria_user):
    categories = ("keys", "bags", "documents", None, "general")
    places = ("Patio", "Biblioteca", "Aula 4", None, "Cafetería", "Laboratorio", "Auditorio")
    specs = [
        dict(
            day=f"2026-10-{1 + i % 31:02d}",
            category=categories[i % len(categories)],
            location=places[i % len(places)],
            status=("lost", "found", "returned")[i % 3],
        )
        for i in range(200)
    ]
    _add_reports(maria_user, *specs)

    statements = _count_queries(maria, "/dashboard/?mes=2026-10")

    assert len(statements) <= 10, statements


def test_query_count_does_not_grow_with_the_number_of_reports(maria, maria_user):
    _add_reports(maria_user, *_month_of(10))
    small = len(_count_queries(maria, "/dashboard/?mes=2026-10"))

    _add_reports(maria_user, *_month_of(300, location="Otro lugar", category="keys"))
    large = len(_count_queries(maria, "/dashboard/?mes=2026-10"))

    assert small == large


def test_aggregations_run_in_sql_not_in_python(maria, maria_user):
    _add_reports(maria_user, *_month_of(30))

    statements = " ".join(_count_queries(maria, "/dashboard/?mes=2026-10")).lower()

    assert statements.count("count(") >= 4
    assert "group by" in statements
    assert "limit" in statements


# --- summary function ---------------------------------------------------------


def test_month_summary_returns_the_documented_metrics(app, maria_user):
    _add_reports(
        maria_user,
        dict(day="2026-10-05", status="returned"),
        dict(day="2026-10-06", status="found"),
        dict(day="2026-09-06"),
    )

    with app.test_request_context("/"):
        summary = dashboard_module.month_summary(2026, 10)

    kpis = summary["kpis"]
    assert (kpis["total"], kpis["lost"], kpis["found"], kpis["returned"]) == (2, 0, 1, 1)
    assert kpis["return_rate"] == "50,0 %"
    assert kpis["delta_vs_previous"] == "+1 respecto al mes anterior"
    assert summary["mes"] == "2026-10"
    assert summary["month_label"] == "octubre de 2026"
    assert summary["prev_mes"] == "2026-09" and summary["next_mes"] == "2026-11"
