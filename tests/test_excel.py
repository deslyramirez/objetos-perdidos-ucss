"""Excel export and template, .xlsx import and the deletion requests the administrator answers."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, inspect

from admin import ensure_moderation_columns
from extensions import db
from models import Conversation, Item, Message, User

HEADER_ROW = 4


def _text(response):
    return response.get_data(as_text=True)


@pytest.fixture
def admin(browser_for):
    browser = browser_for("jefa")
    User.query.filter_by(username="jefa").one().is_admin = True
    db.session.commit()
    return browser


@pytest.fixture
def carlos(admin, browser_for):
    return browser_for("carlos")


def _user(name):
    return User.query.filter_by(username=name).one()


def _workbook(response):
    return load_workbook(BytesIO(response.get_data()))


def _table(sheet):
    """Data rows of the «Reportes» sheet as dicts keyed by the header titles."""
    rows = list(sheet.iter_rows(min_row=HEADER_ROW, values_only=True))
    header = rows[0]
    return [dict(zip(header, row)) for row in rows[1:] if any(cell is not None for cell in row)]


def _upload(browser, content, filename="reportes.xlsx"):
    return browser.post(
        "/import/csv",
        data={"file": (BytesIO(content), filename)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )


# --- export -------------------------------------------------------------------


def test_export_is_a_formatted_workbook_with_a_summary(admin, carlos, make_item):
    make_item(_user("carlos"), name="Llaves", category="keys", status="found", event_date=date(2026, 3, 2))
    make_item(_user("carlos"), name="=SUMA(1)", event_date=date(2026, 3, 1))

    response = admin.get("/export/excel")

    assert response.status_code == 200
    assert response.headers["Content-Disposition"] == "attachment; filename=objetos.xlsx"
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = _workbook(response)
    assert wb.sheetnames == ["Reportes", "Resumen"]
    sheet = wb["Reportes"]
    assert sheet["A1"].value == "Reporte de objetos perdidos y encontrados"
    assert "2 reportes" in sheet["A2"].value
    header = sheet[HEADER_ROW]
    assert header[0].font.bold and header[0].fill.fgColor.rgb.endswith("0D7180")
    assert sheet.freeze_panes == "D5"
    assert sheet.auto_filter.ref == "A4:M6"

    rows = _table(sheet)
    assert [row["Objeto"] for row in rows] == ["=SUMA(1)", "Llaves"]  # by event date
    assert sheet["C5"].data_type == "s"  # a name never runs as a formula
    assert rows[1]["Categoría"] == "Llaves" and rows[1]["Estado"] == "Encontrado"
    assert rows[1]["Publicación"] == "Publicado" and rows[1]["Reportado por"] == "carlos"
    assert rows[1]["Fecha del suceso"].date() == date(2026, 3, 2)
    assert sheet["E5"].number_format == "DD/MM/YYYY"

    summary = [row for row in wb["Resumen"].iter_rows(values_only=True) if row[1] is not None]
    assert ("Perdido", 1) in [(row[1], row[2]) for row in summary]
    assert ("Encontrado", 1) in [(row[1], row[2]) for row in summary]


def test_export_keeps_deactivated_reports_and_drops_unpublished_ones(admin, carlos, make_item):
    make_item(_user("carlos"), name="Visible")
    make_item(_user("carlos"), name="Oculto", approval="inactive")
    make_item(_user("carlos"), name="Pendiente", approval="pending")
    make_item(_user("carlos"), name="Rechazado", approval="rejected")

    rows = _table(_workbook(admin.get("/export/excel"))["Reportes"])

    assert {row["Objeto"]: row["Publicación"] for row in rows} == {"Visible": "Publicado", "Oculto": "Desactivado"}


def test_export_filters_by_month(admin, carlos, make_item):
    make_item(_user("carlos"), name="Octubre", event_date=date(2026, 10, 5))
    make_item(_user("carlos"), name="Septiembre", event_date=date(2026, 9, 5))

    response = admin.get("/export/excel?mes=2026-10")

    assert response.headers["Content-Disposition"] == "attachment; filename=objetos_2026-10.xlsx"
    sheet = _workbook(response)["Reportes"]
    assert [row["Objeto"] for row in _table(sheet)] == ["Octubre"]
    assert "Octubre de 2026" in sheet["A2"].value.lower().capitalize() or "octubre de 2026" in sheet["A2"].value.lower()


def test_export_and_template_are_for_the_administrator_only(carlos):
    for path in ("/export/excel", "/import/plantilla"):
        response = carlos.get(path)
        assert response.status_code == 302, path


# --- template -----------------------------------------------------------------


def test_template_is_a_clean_sheet_with_titles_and_dropdowns(admin):
    response = admin.get("/import/plantilla")

    assert response.headers["Content-Disposition"] == "attachment; filename=plantilla_importar_reportes.xlsx"
    wb = _workbook(response)
    # The instructions live on the Importar page, not in the file.
    assert wb.sheetnames == ["Reportes"]
    sheet = wb["Reportes"]
    titles = [cell.value for cell in sheet[HEADER_ROW]]
    assert titles == ["Objeto", "Lugar", "Fecha del suceso", "Descripción", "Categoría", "Estado", "Contacto"]
    formulas = [validation.formula1 for validation in sheet.data_validations.dataValidation]
    assert any("Electrónica" in (formula or "") for formula in formulas)
    assert any("Perdido,Encontrado,Entregado" in (formula or "") for formula in formulas)


def test_a_filled_template_imports(admin):
    wb = load_workbook(BytesIO(admin.get("/import/plantilla").get_data()))
    sheet = wb["Reportes"]
    sheet["A5"], sheet["B5"], sheet["C5"] = "Mochila negra", "Aula 101", date(2026, 3, 4)
    sheet["D5"], sheet["E5"], sheet["F5"], sheet["G5"] = "Con cuaderno", "Mochilas y útiles", "Encontrado", "999 000 111"
    sheet["A6"], sheet["B6"], sheet["C6"] = "Sin fecha", "Patio", None
    sheet["D6"], sheet["E6"], sheet["F6"], sheet["G6"] = "Azul", "Otros", "Perdido", "999"
    sheet["A7"], sheet["B7"], sheet["C7"] = "Sin contacto", "Patio", date(2026, 3, 4)
    sheet["D7"], sheet["E7"], sheet["F7"] = "Azul", "Otros", "Perdido"
    buffer = BytesIO()
    wb.save(buffer)

    html = _text(_upload(admin, buffer.getvalue()))

    assert "Se importó 1 reporte." in html
    assert "fila 6 (fecha inválida)" in html and "fila 7 (contacto requerido)" in html
    item = Item.query.one()
    assert (item.name, item.location, item.event_date) == ("Mochila negra", "Aula 101", date(2026, 3, 4))
    assert item.description == "Con cuaderno"
    assert (item.category, item.status, item.contact) == ("bags", "found", "999 000 111")
    assert item.approval == "approved"


def test_an_exported_file_can_be_imported_again(admin, carlos, make_item):
    make_item(_user("carlos"), name="Paraguas", category="other", status="returned", description="Negro", contact="999")
    exported = admin.get("/export/excel").get_data()
    Item.query.delete()
    db.session.commit()

    html = _text(_upload(admin, exported, "objetos.xlsx"))

    assert "Se importó 1 reporte." in html
    item = Item.query.one()
    assert (item.name, item.category, item.status) == ("Paraguas", "other", "returned")


def test_a_plain_workbook_with_spanish_titles_imports(admin):
    wb = Workbook()
    wb.active.append(["objeto", "LUGAR", "Fecha", "descripcion", "Categoria", "ESTADO", "contacto"])
    wb.active.append(["Celular", "Biblioteca", "15/03/2026", "Funda roja", "Electrónica", "Perdido", "999"])
    buffer = BytesIO()
    wb.save(buffer)

    assert "Se importó 1 reporte." in _text(_upload(admin, buffer.getvalue()))


def test_a_broken_xlsx_is_rejected(admin):
    html = _text(_upload(admin, b"PK\x03\x04 esto no es un excel"))

    assert "Verifique que sea un Excel (.xlsx) válido." in html
    assert Item.query.count() == 0


def test_a_csv_with_spanish_titles_imports(admin):
    content = (
        "Objeto;Lugar;Fecha del suceso;Descripción;Categoría;Estado;Contacto\n"
        "Llaves;Patio;2026-03-02;Tres llaves;Llaves;Encontrado;999\n"
    ).encode("utf-8")

    assert "Se importó 1 reporte." in _text(_upload(admin, content, "datos.csv"))
    assert Item.query.one().category == "keys"


# --- deletion requests ----------------------------------------------------------


def test_the_request_reaches_the_moderation_queue(admin, carlos, make_item):
    item = make_item(_user("carlos"), name="Billetera")
    carlos.post(f"/delete/{item.id}", data={"reason": "La publiqué dos veces"})

    page = _text(admin.get("/admin/moderacion?tab=eliminacion"))

    assert "Billetera" in page and "La publiqué dos veces" in page
    assert "Desactivar" in page and "Eliminar definitivamente" in page and "Mantener publicado" in page
    assert 'class="nav-badge"' in page  # counted in the menu badge


def test_deactivating_hides_the_report_but_keeps_it_in_the_excel(admin, carlos, browser_for, make_item):
    item = make_item(_user("carlos"), name="Billetera")
    carlos.post(f"/delete/{item.id}")

    admin.post(f"/admin/item/{item.id}/desactivar")

    item = db.session.get(Item, item.id)
    assert item.approval == "inactive" and item.deletion_requested_at is None
    other = browser_for("ana")
    assert "Billetera" not in _text(other.get("/"))
    assert other.get(f"/item/{item.id}").status_code == 404
    # The author no longer sees it either: only the administrator does.
    assert "Billetera" not in _text(carlos.get("/"))
    assert carlos.get(f"/item/{item.id}").status_code == 404
    assert admin.get(f"/item/{item.id}").status_code == 200
    assert "Billetera" in _text(admin.get("/admin/moderacion?tab=desactivadas"))
    rows = _table(_workbook(admin.get("/export/excel"))["Reportes"])
    assert [(row["Objeto"], row["Publicación"]) for row in rows] == [("Billetera", "Desactivado")]


def test_reactivating_publishes_again(admin, carlos, make_item):
    item = make_item(_user("carlos"), name="Gorra", approval="inactive")

    admin.post(f"/admin/item/{item.id}/reactivar")

    assert db.session.get(Item, item.id).approval == "approved"


def test_keeping_the_report_clears_the_request(admin, carlos, make_item):
    item = make_item(_user("carlos"), name="Gorra")
    carlos.post(f"/delete/{item.id}", data={"reason": "x"})

    admin.post(f"/admin/item/{item.id}/mantener")

    item = db.session.get(Item, item.id)
    assert item.approval == "approved" and item.deletion_requested_at is None and item.deletion_reason is None


def test_deleting_for_good_removes_the_report_its_messages_and_its_excel_row(admin, carlos, browser_for, make_item):
    item = make_item(_user("carlos"), name="Audífonos")
    ana = browser_for("ana")
    ana.post(f"/messages/start/{item.id}", data={"kind": "owner", "body": "Creo que son míos"})
    assert Conversation.query.count() == 1
    carlos.post(f"/delete/{item.id}")

    response = admin.post(
        f"/admin/item/{item.id}/eliminar", data={"next": f"/item/{item.id}"}, follow_redirects=True
    )

    assert response.request.path == "/admin/moderacion"
    assert "se eliminó definitivamente" in _text(response)
    assert Item.query.count() == 0 and Conversation.query.count() == 0 and Message.query.count() == 0
    sheet = _workbook(admin.get("/export/excel"))["Reportes"]
    assert [row["Objeto"] for row in _table(sheet) if row["Objeto"]] == []
    assert sheet["A5"].value == "No hay reportes para los filtros elegidos."


def test_regular_users_cannot_answer_requests(carlos, make_item):
    item = make_item(_user("carlos"), name="Gorra")
    for action in ("desactivar", "reactivar", "mantener", "eliminar"):
        carlos.post(f"/admin/item/{item.id}/{action}")

    assert db.session.get(Item, item.id).approval == "approved"


def test_the_deletion_columns_are_added_to_old_databases(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE user (id INTEGER PRIMARY KEY, username VARCHAR(64))")
        conn.exec_driver_sql("CREATE TABLE item (id INTEGER PRIMARY KEY, name VARCHAR(120))")
    ensure_moderation_columns(engine)
    ensure_moderation_columns(engine)
    columns = {c["name"] for c in inspect(engine).get_columns("item")}
    engine.dispose()

    assert {"deletion_requested_at", "deletion_reason"} <= columns
